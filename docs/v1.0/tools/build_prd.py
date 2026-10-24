from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
VERSION=(ROOT/'VERSION').read_text().strip()
doc=Document();s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.75);s.bottom_margin=Inches(.7);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
s.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',25),('Subtitle',11),('Heading 1',16),('Heading 2',12.5)]:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 st.paragraph_format.space_after=Pt(7)
 if name.startswith('Heading'):
  st.paragraph_format.space_before=Pt(15 if name=='Heading 1' else 10);st.paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.line_spacing=1.22
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Title'].paragraph_format.space_after=Pt(14)
def table_style(t):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(t.columns);widths=[.8,2.1,2.4,1.6] if n==4 else ([1.2,2.9,2.8] if n==3 else [1.6,5.3])
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   tcpr.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'100');x.set(qn('w:type'),'dxa');m.append(x)
   tcpr.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'344A40' if ri==0 else ('F3F5F2' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.12
    if ci==0 and n==4:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '26382C');r.bold=ri==0
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
lines=(ROOT/'拾史桌面端V1_PRD.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   parts=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',x) for x in parts):rows.append(parts)
   i+=1
  t=doc.add_table(rows=0,cols=len(rows[0]))
  for row in rows:
   for c,txt in zip(t.add_row().cells,row):c.text=txt
  table_style(t);doc.add_paragraph().paragraph_format.space_after=Pt(2);continue
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('- '):
  p=doc.add_paragraph('• '+line[2:]);p.paragraph_format.left_indent=Inches(.1);p.paragraph_format.space_after=Pt(4)
 else:
  items=re.split(r'(?=F\d{2}：)',line)
  if len(items)>1:
   for item in items:
    if not item.strip():continue
    p=doc.add_paragraph(item.strip());p.paragraph_format.space_after=Pt(5)
    m=re.match(r'(F\d{2}：)(.*)',item.strip())
    if m:p.clear();p.add_run(m[1]).bold=True;p.add_run(m[2])
  else:doc.add_paragraph(line)
 i+=1
p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('拾史 v'+VERSION+'  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('7B8279')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
for st in doc.styles:
 for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
for b in list(doc.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.core_properties.title='拾史文化故事视频工作台产品需求文档';doc.core_properties.author=''
dest=ROOT/'拾史桌面端V1_PRD.docx';doc.save(dest)
(ROOT/'build').mkdir(exist_ok=True)
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
(ROOT/'build/prd.json').write_text(json.dumps({'version':VERSION,'sources':{'拾史桌面端V1_PRD.md':digest(ROOT/'拾史桌面端V1_PRD.md'),'tools/build_prd.py':digest(Path(__file__))},'outputs':{'拾史桌面端V1_PRD.docx':digest(dest)}},ensure_ascii=False,indent=2),encoding='utf-8')
print(dest);print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
